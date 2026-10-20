"""Журнал ставок -> Excel (с фильтрами по колонкам)."""
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

EVENTS = {
    "initial_sent": "Первая ставка отправлена",
    "initial_prepared": "Первая ставка (тренировка, не отправлена)",
    "bid_sent": "Повышение ставки отправлено",
    "bid_prepared": "Повышение (тренировка, не отправлено)",
    "bid": "Решено поднять ставку",
    "hold": "Ставку не меняем",
    "over_cap": "Нужно больше предела, не ставим",
    "skipped": "Пропущено (подать нельзя)",
    "finished": "Приём закончен",
    "error": "Ошибка",
}


def journal_to_xlsx(csv_path: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Журнал"
    ws.append(["Время", "ИНН должника", "Номер дела", "Должник", "Событие", "Код события", "Баллы", "Подробности"])
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            ev = r.get("event", "")
            pts = r.get("points", "")
            try:
                pts = float(pts) if pts != "" else ""
            except ValueError:
                pass
            ws.append([r.get("time", "").replace("T", " "), str(r.get("inn", "")),
                       r.get("case", ""), r.get("debtor", ""), EVENTS.get(ev, ev), ev, pts, r.get("detail", "")])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, w in enumerate([19, 16, 18, 34, 38, 16, 9, 90], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "@"                 # ИНН как текст, ведущие нули не теряются
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
