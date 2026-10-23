"""Сводка по протоколам: итог каждого дела одной строкой (кто выиграл, прошли ли мы, ставки)."""
from __future__ import annotations

import io
import re
import warnings
from decimal import Decimal, ROUND_CEILING
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .calc import d, effective_bids, min_bid, winners, CENT
from .protocol import read_protocol

OUR_INN = "7730309884"


def _rows(path: str) -> list[list[str]]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(path).active
    out = []
    for r in ws.iter_rows(values_only=True):
        cells = [("" if c is None else str(c).strip()) for c in r]
        cells = [c for c in cells if c]
        if cells:
            out.append(cells)
    return out


def _after(rows, key: str, start: int = 0) -> str:
    """Значение справа от подписи в строке вида [подпись, значение]."""
    for r in rows[start:]:
        if r[0] == key and len(r) > 1:
            return r[1]
    return ""


def min_to_pass(rivals: list[Decimal]) -> Decimal:
    """Наименьший балл, при котором наша СРО проходит по среднему (не обязательно единственной)."""
    if not rivals:
        return CENT
    x = (sum(rivals) / len(rivals)).quantize(CENT, rounding=ROUND_CEILING)
    while True:
        bids = {str(i): v for i, v in enumerate(rivals)}
        bids["our"] = x
        if "our" in winners(bids):
            return x
        x += CENT


def summarize(path: str, our_inn: str = OUR_INN) -> dict:
    rows = _rows(path)
    p = read_protocol(path)
    bids = effective_bids(p.applications)
    ours = bids.get(our_inn)
    rivals = [v for k, v in bids.items() if k != our_inn]
    top = max(rivals) if rivals else None

    win_sro = win_au = ""
    for i, r in enumerate(rows):
        if r[0].startswith("Кандидатура") and "выбран" in r[0] and len(r) == 1:
            win_sro_inn = _after(rows, "ИНН СРО", i)
            win_sro = _after(rows, "Наименование СРО", i)
            win_au = _after(rows, "Ф.И.О. АУ", i)
            break
    else:
        win_sro_inn = ""
    won = bool(win_sro_inn) and win_sro_inn == our_inn

    res = {
        "end": _after(rows, "Дата окончания процедуры выбора"),
        "inn": _after(rows, "ИНН должника"),
        "debtor": p.debtor,
        "case": next((r[1] for r in rows if r[0] == "Номер дела" and len(r) > 1), ""),
        "n_sro": len(bids),
        "ours": ours,
        "top": top,
        "avg": p.average,
        "passed": ("нет заявки" if ours is None else ("да" if our_inn in p.selected else "нет")),
        "win_sro": win_sro,
        "win_au": win_au,
        "won": "нет заявки" if ours is None else ("ДА" if won else "нет"),
        "need_sole": min_bid(rivals).points if rivals else None,
        "need_pass": min_to_pass(rivals) if rivals else None,
        "file": Path(path).name,
    }
    return res


HEAD = [
    ("end", "Окончание приёма", 14), ("inn", "ИНН должника", 14), ("debtor", "Должник", 38),
    ("case", "Номер дела", 16), ("n_sro", "СРО в приёме", 9), ("ours", "Наша ставка", 10),
    ("top", "Макс. ставка конкурента", 12), ("avg", "Средний балл", 10),
    ("passed", "Прошли по баллу", 11), ("won", "Мы победили", 11),
    ("win_sro", "Победившая СРО", 36), ("win_au", "Победивший АУ", 30),
    ("need_pass", "Минимум, чтобы пройти по баллу", 14),
    ("need_sole", "Минимум, чтобы остальные не прошли", 14), ("file", "Файл протокола", 44),
]


def build_xlsx(folder: str, our_inn: str = OUR_INN) -> tuple[bytes, list[str]]:
    files = sorted(Path(folder).rglob("*.xlsx"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.append([h[1] for h in HEAD])
    errors: list[str] = []
    for f in files:
        try:
            s = summarize(str(f), our_inn)
        except Exception as e:                       # протокол другого вида или битый файл
            errors.append(f"{f.name}: {e!r}")
            continue
        ws.append([float(s[k]) if isinstance(s[k], Decimal) else (s[k] if s[k] is not None else "")
                   for k, _, _ in HEAD])
        row = ws.max_row
        ws.cell(row, 2).number_format = "@"
        if s["won"] == "ДА":
            for c in ws[row]:
                c.fill = PatternFill("solid", fgColor="E2F0D9")
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, (_, _, w) in enumerate(HEAD, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if errors:
        e = wb.create_sheet("Не разобрано")
        e.append(["Файл и причина"])
        for x in errors:
            e.append([x])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), errors
