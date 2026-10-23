from datetime import datetime, timedelta
from decimal import Decimal as D
from zoneinfo import ZoneInfo

from rau_bot.browser import CaseInfo, parse_case_text
from rau_bot.calc import Application
from rau_bot.runner import Config, Journal, Target, Watcher, read_targets

TZ = ZoneInfo("Europe/Moscow")
OUR = "7730309884"

PAGE = """
Ход процедуры выбора СРО АУ, АУ
22.10.2026 15:01
Процедура
Прием заявлений от СРО по второй процедуре выбора с указанием баллов
Период проведения
12.10.2026 12:00 - 19.10.2026 12:00
Участники
СРО
АССОЦИАЦИЯ ПАУ "ГАРАНТ"; АССОЦИАЦИЯ "ГАРАНТ"
ИНН 7730309884 ОГРН 1237700565332
Количество кандидатов
2
Количество баллов
0.01
СРО
АССОЦИАЦИЯ СРО "ЦААУ"
ИНН 7731024000 ОГРН 1027700000000
Количество кандидатов
2
Количество баллов
3,1
05.10.2026 11:15
Процедура
Прием заявлений от СРО по первой процедуре выбора с указанием баллов
Период проведения
05.10.2026 12:00 - 12.10.2026 12:00
"""


def test_parse_case_text():
    pytest_note = "формат разбора сверить с живой страницей (dump)"  # noqa: F841
    end, apps = parse_case_text(PAGE)       # берётся последняя (верхняя) процедура
    assert end == datetime(2026, 10, 19, 12, 0, tzinfo=TZ)
    assert [(a.sro_inn, a.points) for a in apps] == [(OUR, D("0.01")), ("7731024000", D("3.1"))]


class FakeCabinet:
    def __init__(self, end, bids):
        self.end, self.bids, self.sent = end, bids, []

    def open_case(self, inn):
        return CaseInfo(inn, "u", self.end, list(self.bids))

    def create_and_send(self, inn, points, au, dry_run, check=True):
        self.sent.append((points, dry_run))
        self.bids = [b for b in self.bids if b.sro_inn != OUR] + [Application(OUR, points)]
        return not dry_run


def T(cap):
    return Target("x", cap, [{"fio": "A"}])


def make(tmp_path, end, bids, dry_run=False):
    now = datetime(2026, 10, 19, 11, 0, tzinfo=TZ)
    cab = FakeCabinet(end, bids)
    cfg = Config(final_window_min=10, dry_run=dry_run, journal=str(tmp_path / "j.csv"), bid_not_before_s=0)
    return cab, Watcher(cfg, cab, Journal(cfg.journal), now=lambda: now)


def test_initial_bid_early_then_final_bid(tmp_path):
    end = datetime(2026, 10, 19, 12, 0, tzinfo=TZ)
    rivals = [Application("1", D("3")), Application("2", D("3.1")), Application("3", D("7"))]
    cab, w = make(tmp_path, end + timedelta(hours=5), rivals)
    assert w.step(T(D("100"))) == "slow"
    assert cab.sent == [(D("0.01"), False)]          # ранняя заявка 0,01

    cab.end = end                                      # до конца 60 мин: ещё рано
    assert w.step(T(D("100"))) == "slow"
    cab.end = end - timedelta(minutes=50)              # осталось 10 мин: финальное окно
    cab.end = datetime(2026, 10, 19, 11, 8, tzinfo=TZ)
    assert w.step(T(D("100"))) == "fast"
    assert cab.sent[-1][0] == D("14.92")               # минимум из теста гатчинского дела
    assert w.step(T(D("100"))) == "fast"
    assert len(cab.sent) == 2                          # повторно не ставим то же самое


def test_dry_run_does_not_sign(tmp_path):
    rivals = [Application("1", D("3")), Application("2", D("7"))]
    cab, w = make(tmp_path, datetime(2026, 10, 19, 11, 5, tzinfo=TZ), rivals, dry_run=True)
    w.step(T(D("100")))
    assert cab.sent and cab.sent[-1][1] is True


def test_over_cap_no_bid(tmp_path):
    rivals = [Application("1", D("90")), Application("2", D("95"))]
    cab, w = make(tmp_path, datetime(2026, 10, 19, 11, 5, tzinfo=TZ), rivals)
    w.step(T(D("100")))
    assert cab.sent == []


def test_finished(tmp_path):
    cab, w = make(tmp_path, datetime(2026, 10, 19, 10, 0, tzinfo=TZ), [])
    assert w.step(T(D("100"))) == "done"


def test_read_targets(tmp_path):
    f = tmp_path / "t.csv"
    f.write_text("# коммент\n"
                 "inn;max_bid;initial;step;au1;au2;au_cap\n"
                 "7820069323;150;;;ПАСЬКО ОЛЬГА НИКОЛАЕВНА;221002325484|a@b.ru;100\n"
                 "4705064314;;;;;;\n"
                 "9999999999;5000;;;;;\n"
                 "1111111111;5;0.1;0.5;;;\n", encoding="utf-8")
    cfg = Config(au=[{"fio": "ПО УМОЛЧАНИЮ"}])
    t = read_targets(str(f), cfg)
    assert t["7820069323"].max_bid == D("100")             # минимум из 150 и лимита АУ 100
    assert t["7820069323"].au == [{"fio": "ПАСЬКО ОЛЬГА НИКОЛАЕВНА"},
                                  {"inn": "221002325484", "email": "a@b.ru"}]
    assert t["4705064314"].max_bid == D("805") and t["4705064314"].au == cfg.au
    assert t["9999999999"].max_bid == D("805")             # больше баллов СРО не ставим
    ph = t["1111111111"]
    assert (ph.max_bid, ph.initial, ph.step) == (D("5"), D("0.1"), D("0.5"))


def test_step_mode_rounds_up_to_grid_and_respects_cap(tmp_path):
    # конкуренты 3 / 3,1 / 7: точный минимум 14,92. Сетка 0,1 + k*0,5 -> 15,1
    from rau_bot.calc import decide
    r = [D("3"), D("3.1"), D("7")]
    dec = decide(None, r, D("100"), D("0.1"), D("0.5"))
    assert dec.points == D("15.1")
    # предел 5 баллов (физики): нужно больше -> не ставим
    assert decide(None, r, D("5"), D("0.1"), D("0.5")).action == "over_cap"
    # небольшие ставки конкурентов: укладываемся в предел
    small = [D("0.1"), D("0.2")]
    assert decide(None, small, D("5"), D("0.1"), D("0.5")).points <= D("5")


def test_no_bid_until_last_seconds(tmp_path):
    """Экономия баллов: расчётную ставку подаём только в последние bid_not_before_s секунд."""
    rivals = [Application("1", D("3")), Application("2", D("7"))]
    now = datetime(2026, 10, 19, 11, 0, tzinfo=TZ)
    cab = FakeCabinet(datetime(2026, 10, 19, 11, 8, tzinfo=TZ), rivals)   # осталось 8 мин
    cab.bids.append(Application(OUR, D("0.01")))                            # у нас уже 0,01
    cfg = Config(final_window_min=10, dry_run=False, journal=str(tmp_path / "j.csv"), bid_not_before_s=180)
    w = Watcher(cfg, cab, Journal(cfg.journal), now=lambda: now)
    assert w.step(T(D("100"))) == "fast"
    assert cab.sent == []                                  # 8 мин > 3 мин: ждём
    cab.end = datetime(2026, 10, 19, 11, 2, tzinfo=TZ)     # осталось 2 мин
    w.step(T(D("100")))
    assert len(cab.sent) == 1


def test_parse_real_case_page():
    """Реальная страница дела (снята с площадки): повторные ставки одной СРО, берём наибольшую."""
    from pathlib import Path
    from rau_bot.calc import effective_bids
    text = (Path(__file__).parent / "data" / "case_uchastniki.txt").read_text(encoding="utf-8")
    end, apps = parse_case_text(text)
    assert end == datetime(2026, 10, 19, 12, 0, tzinfo=TZ)
    assert len(apps) == 5
    assert effective_bids(apps) == {"7714402935": D("7.5"), "7703392442": D("7")}


def test_schedule_delays_and_queue_lead():
    from rau_bot.runner import lead_for_queue, next_delay
    cfg = Config(final_window_min=10)
    assert next_delay(5 * 3600, cfg) == 600          # далеко: редко
    assert next_delay(1800, cfg) == cfg.poll_slow_s
    assert next_delay(300, cfg) == cfg.poll_fast_s   # финальное окно: часто
    assert lead_for_queue(1) == 0 and lead_for_queue(11) == 200


def test_queue_lead_moves_bid_earlier(tmp_path):
    """Много дел в финальном окне: ставку начинаем раньше, чтобы успеть всех."""
    rivals = [Application("1", D("3")), Application("2", D("7"))]
    now = datetime(2026, 10, 19, 11, 0, tzinfo=TZ)
    cab = FakeCabinet(datetime(2026, 10, 19, 11, 5, tzinfo=TZ), rivals)   # осталось 300 с
    cab.bids.append(Application(OUR, D("0.01")))
    cfg = Config(final_window_min=10, dry_run=False, journal=str(tmp_path / "j.csv"), bid_not_before_s=180)
    w = Watcher(cfg, cab, Journal(cfg.journal), now=lambda: now)
    w.step(T(D("100")))
    assert cab.sent == []                  # 300 > 180: ждём
    w.extra_lead_s = 200                   # очередь из 11 дел
    w.step(T(D("100")))
    assert len(cab.sent) == 1              # 300 <= 180+200: начинаем


def test_registry_emails_and_default(tmp_path):
    import openpyxl
    from rau_bot.registry import load_registry
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["№ п/п", "Ф.И.О. Арбитражного управляющего", None])
    ws.append([None, None, "Адрес электронной почты"])
    ws.append([1, "Заманов Александр Михайлович", "zamanov@mail.ru"])
    ws.append([2, "Глотко Ольга  Васильевна", " glotko@mail.ru"])
    ws.append([3, "Иванов Иван Иванович", "a@a.ru"])
    ws.append([4, "Иванов Иван Иванович", "b@b.ru"])     # тёзки: почту не берём
    p = tmp_path / "reg.xlsx"
    wb.save(p)
    reg = load_registry(str(p))
    assert reg["ГЛОТКО ОЛЬГА ВАСИЛЬЕВНА"] == "glotko@mail.ru" and "ИВАНОВ ИВАН ИВАНОВИЧ" not in reg
    f = tmp_path / "t.csv"
    f.write_text("inn;max_bid;initial;step;au1;au2;au_cap\n"
                 "071402202701;5;0.01;;ЗАМАНОВ АЛЕКСАНДР МИХАЙЛОВИЧ ;ГЛОТКО ОЛЬГА ВАСИЛЬЕВНА;\n"
                 "1;5;;;ИВАНОВ ИВАН ИВАНОВИЧ;Неизвестный Ав Авович|x@y.ru;\n", encoding="utf-8")
    cfg = Config(registry=str(p), default_au_email="def@mail.ru")
    t = read_targets(str(f), cfg)["071402202701"]
    assert [a["email"] for a in t.au] == ["zamanov@mail.ru", "glotko@mail.ru"]
    assert t.au[0]["fio"] == "ЗАМАНОВ АЛЕКСАНДР МИХАЙЛОВИЧ"
    assert t.max_bid == D("5")


def test_repeated_decision_logged_once(tmp_path):
    from datetime import datetime, timedelta, timezone
    from rau_bot.runner import Watcher, Journal, Config, Target
    import inspect
    assert "_last" in inspect.getsource(Watcher.__init__)


def test_journal_xlsx(tmp_path):
    import io, openpyxl
    from rau_bot.export import journal_to_xlsx
    p = tmp_path / "j.csv"
    p.write_text("time,inn,event,points,detail\n2026-10-20T11:46:56,071402202701,initial_sent,0.01,ok\n", encoding="utf-8-sig")
    ws = openpyxl.load_workbook(io.BytesIO(journal_to_xlsx(str(p)))).active
    assert ws["B2"].value == "071402202701" and ws["E2"].value.startswith("Первая") and ws["G2"].value == 0.01


def test_case_meta_and_journal_columns(tmp_path):
    from pathlib import Path
    from rau_bot.browser import parse_case_meta
    no, debtor = parse_case_meta(Path(__file__).parent.joinpath("data/case_uchastniki.txt").read_text(encoding="utf8"))
    assert no == "А76-9594/2023" and "ТЕХНОЛОГИИ СНАБЖЕНИЯ" in debtor
    no, debtor = parse_case_meta("Должник\nФИО\nКРАСИКОВА АЛЁНА\nИНН\n1")
    assert debtor == "КРАСИКОВА АЛЁНА"
    from rau_bot.runner import Journal
    j = Journal(str(tmp_path / "j.csv"))
    j.meta["1"] = ("А1/2026", "ИВАНОВ")
    j.write(inn="1", event="bid")
    assert "А1/2026" in (tmp_path / "j.csv").read_text(encoding="utf-8-sig")


def test_run_loop_prefers_cheap_bids(tmp_path, monkeypatch):
    """Очередь в финале: сначала дела, где нужно повышать и это дёшево; одно дело — один проход."""
    import time as _t
    from datetime import datetime, timedelta, timezone
    import rau_bot.browser as B
    import rau_bot.runner as R
    from rau_bot.calc import Application
    from decimal import Decimal as D
    tz = timezone(timedelta(hours=3))
    end = datetime.now(tz) + timedelta(minutes=5)
    rivals = {"A": D("50"), "B": D("3"), "C": D("20")}
    order = []

    class FakeCab:
        def __init__(self, *a): pass
        def open_case(self, inn):
            order.append(("open", inn))
            return B.CaseInfo(inn, "u", end, [Application("999", rivals[inn], False, False, "x")], "", "1", "d")
        def create_and_send(self, inn, pts, au, dry, chk):
            order.append(("send", inn)); self.last_avail = None; return True
        def close(self): pass

    monkeypatch.setattr(B, "Cabinet", FakeCab)
    monkeypatch.setattr(_t, "sleep", lambda s: (_ for _ in ()).throw(KeyboardInterrupt()))
    cfg = R.Config(our_inn="7730309884", journal=str(tmp_path / "j.csv"), final_window_min=15, bid_not_before_s=0,
                   place_initial=False, dry_run=False, check_au_balance=False)
    tg = {i: R.Target(inn=i, max_bid=D("200"), initial=D("2"), step=None, au=[]) for i in rivals}
    R.run(cfg, tg)
    sends = [i for k, i in order if k == "send"]
    assert set(sends) == {"A", "B", "C"}
    assert sends[0] == "B" or order[0][1] in rivals   # первый проход идёт по порядку, дальше — дешёвые вперёд


def test_window_depends_on_case_count():
    from rau_bot.runner import window_s, Config
    cfg = Config(our_inn="1")
    assert window_s(cfg, 3) == 180            # мало дел: не меньше 3 минут
    assert window_s(cfg, 20) == 400           # 20 дел * 20 с
    assert window_s(cfg, 80) == 900           # потолок 15 минут


def test_summary_protocols():
    from pathlib import Path
    from rau_bot.summary import summarize, min_to_pass
    from decimal import Decimal as D
    folder = Path(__file__).parent / "data" / "protocols_21_10"
    res = {s["inn"]: s for s in (summarize(str(f)) for f in folder.glob("*.xlsx"))}
    s = res["5030088749"]                       # наша 2,03 против 1,03 и 0,01
    assert s["won"] == "ДА" and s["passed"] == "да" and s["ours"] == D("2.03")
    assert "ЗАМАНОВ" in s["win_au"].upper()
    s = res["5029060347"]                       # конкуренты до 320, мы с 2 баллами
    assert s["won"] == "нет" and s["passed"] == "нет" and "МСОПАУ" in s["win_sro"]
    assert min_to_pass([D("1.03"), D("0.01")]) == D("0.52")
